# -*- coding: utf-8 -*-
"""Fügt in die vom Nutzer bearbeitete Slide 6 die Phase 'Besuchstermin'
zwischen DD (3) und Verbindl. Angebot ein. Klont vorhandene Elemente,
damit Stil 1:1 erhalten bleibt. Verändert sonst nichts."""
import copy
from pptx import Presentation
from pptx.util import Inches, Emu

SRC="user_ip.pptx"; OUT="user_ip_besuch.pptx"
pr=Presentation(SRC)
sl=pr.slides[5]
sh=list(sl.shapes)
spTree=sl.shapes._spTree

# --- Indizes der 7 vorhandenen Phasen (links->rechts) ---
BOX =[13,16,19,22,25,28,31]
NUM =[14,17,20,23,26,29,32]
NAME=[15,18,21,24,27,30,33]
DIA =[34,36,38,40,42,44,46]
LAB =[35,37,39,41,43,45,47]
TICK=[63,64,65,66,67,68,69]
WLAB=[61,62,70,71,72,73,74]

# --- Ziel-Layout: 8 Phasen gleichmäßig über 1.56 .. 12.54 ---
X0=1.56; N=8; GAP=0.08
BOXW=(10.98-(N-1)*GAP)/N          # ~1.3025
STEP=BOXW+GAP
def Lx(pos): return X0+pos*STEP
def Cx(pos): return Lx(pos)+BOXW/2

def setpos(shape,l=None,t=None,w=None,h=None):
    if l is not None: shape.left=Inches(l)
    if t is not None: shape.top=Inches(t)
    if w is not None: shape.width=Inches(w)
    if h is not None: shape.height=Inches(h)

def set_single_text(shape,txt):
    p=shape.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text=txt
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    else:
        p.add_run().text=txt

def clone(src_idx):
    el=copy.deepcopy(sh[src_idx]._element)
    spTree.append(el)
    return sl.shapes[-1]

def in_inch(shape):  # current width in inches
    return Emu(shape.width).inches

# orig Phase k -> neue Position (Besuchstermin = neue Pos 4)
def newpos(k): return k if k<=3 else k+1

# --- 1) vorhandene Phasen umpositionieren ---
for k in range(7):
    pos=newpos(k)
    bx=Lx(pos); c=Cx(pos)
    # Box
    setpos(sh[BOX[k]], l=bx, t=2.74, w=BOXW, h=0.66)
    # Nummer
    setpos(sh[NUM[k]], l=bx+0.06, t=2.76, w=0.4, h=0.3)
    # Name
    setpos(sh[NAME[k]], l=bx+0.05, t=3.06, w=BOXW-0.1, h=0.32)
    # Milestone-Raute (0.48 breit)
    setpos(sh[DIA[k]], l=c-0.24, t=3.58, w=0.48, h=0.48)
    # Milestone-Label (Breite behalten, zentriert, geclamped)
    lw=in_inch(sh[LAB[k]])
    llx=max(1.56,min(c-lw/2,12.54-lw))
    setpos(sh[LAB[k]], l=llx, t=4.08)
    # Wochen-Tick
    setpos(sh[TICK[k]], l=c-0.045, t=2.15)
    # Wochen-Label (Breite behalten, zentriert, geclamped)
    ww=in_inch(sh[WLAB[k]])
    wlx=max(1.56,min(c-ww/2,12.54-ww))
    setpos(sh[WLAB[k]], l=wlx)

# --- 2) Nummern/Milestone-Text der nach rechts gerückten Phasen anpassen ---
for k,newn in [(4,"5"),(5,"6"),(6,"7")]:
    set_single_text(sh[NUM[k]], newn)                 # Phasennummer
    dp=sh[DIA[k]].text_frame.paragraphs[0]
    if len(dp.runs)>=2: dp.runs[1].text=newn          # "M4"->"M5" etc.
    else: set_single_text(sh[DIA[k]], "M"+newn)

# --- 3) neue Phase 4 = Besuchstermin (Klone von Phase 3 / DD) ---
POS=4; bx=Lx(POS); c=Cx(POS)
nb=clone(BOX[3]);  setpos(nb, l=bx, t=2.74, w=BOXW, h=0.66)
nn=clone(NUM[3]);  setpos(nn, l=bx+0.06, t=2.76, w=0.4, h=0.3);  set_single_text(nn,"4")
nm=clone(NAME[3]); setpos(nm, l=bx+0.05, t=3.06, w=BOXW-0.1, h=0.32); set_single_text(nm,"Besuchstermin")
nd=clone(DIA[3]);  setpos(nd, l=c-0.24, t=3.58, w=0.48, h=0.48)
dp=nd.text_frame.paragraphs[0]
if len(dp.runs)>=2: dp.runs[1].text="4"
nl=clone(LAB[3])
nlw=in_inch(nl); nl_l=max(1.56,min(c-nlw/2,12.54-nlw)); setpos(nl, l=nl_l, t=4.08); set_single_text(nl,"Vor Ort bestätigt")
nt=clone(TICK[3]); setpos(nt, l=c-0.045, t=2.15)
nw=clone(WLAB[3])
nww=in_inch(nw); nw_l=max(1.56,min(c-nww/2,12.54-nww)); setpos(nw, l=nw_l); set_single_text(nw,"W 4")

# --- 4) Titel 0–6 -> 0–7 ---
for p in sh[0].text_frame.paragraphs:
    for r in p.runs:
        if "0–6" in r.text: r.text=r.text.replace("0–6","0–7")

pr.save(OUT)
print("SAVED",OUT,"slides:",len(pr.slides._sldIdLst))
