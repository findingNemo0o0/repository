#!/usr/bin/env python3
"""Geometriepruefung: Rahmen, Textueberlauf und Ueberlappungen.

Ersatz fuer die visuelle Kontrolle, solange LibreOffice in dieser Umgebung
nicht laeuft.
"""
import math
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SW, SH = 13.333, 7.5
SAFE_BOTTOM = 7.42
# Faustwert: mittlere Zeichenbreite von Open Sans / Quicksand in em
CHAR_EM = 0.505


def E(v):
    return v / 914400


def para_lines(p, w_in):
    txt = "".join(r.text for r in p.runs)
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    sz = max(sizes) if sizes else 12.0
    if not txt:
        return 1, sz
    cpl = max(int((w_in * 72) / (sz * CHAR_EM)), 6)
    return max(1, math.ceil(len(txt) / cpl)), sz


def est_h(tf, w_in):
    total = 0.0
    for p in tf.paragraphs:
        lines, sz = para_lines(p, w_in)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.10
        total += lines * sz * ls * 1.20 / 72.0
        total += (p.space_before.pt if p.space_before else 0) / 72.0
        total += (p.space_after.pt if p.space_after else 0) / 72.0
    return total


def main(path):
    prs = Presentation(path)
    issues = []
    for i, s in enumerate(prs.slides, 1):
        boxes = []
        cards = []
        for sh in s.shapes:
            if sh.width is None or sh.left is None:
                continue
            x, y, w, h = E(sh.left), E(sh.top), E(sh.width), E(sh.height)
            if x < -0.03 or y < -0.03 or x + w > SW + 0.03 or y + h > SH + 0.03:
                issues.append(f"S{i} ausserhalb der Folie: {sh.name} "
                              f"({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not (
                    sh.has_text_frame and sh.text_frame.text.strip()):
                cards.append((sh.name, x, y, w, h))
            if sh.has_text_frame and sh.text_frame.text.strip():
                eh = est_h(sh.text_frame, w - 0.04)
                boxes.append((sh.name, x, y, w, h, eh, sh.text_frame.text))
                if eh > h + 0.05:
                    issues.append(f"S{i} Textbox zu klein: '{sh.name}' h={h:.2f} "
                                  f"noetig≈{eh:.2f} @({x:.2f},{y:.2f}) w={w:.2f} "
                                  f"→ {sh.text_frame.text[:55]!r}")
                if y + eh > SAFE_BOTTOM:
                    issues.append(f"S{i} laeuft unten heraus: y={y:.2f}+{eh:.2f} "
                                  f"→ {sh.text_frame.text[:55]!r}")
        # Text, der aus seiner Karte herauslaeuft
        for nb, x, y, w, h, eh, txt in boxes:
            for nc, cx, cy, cw, chh in cards:
                inside = (cx - 0.02 <= x and x + w <= cx + cw + 0.02
                          and cy - 0.02 <= y < cy + chh - 0.02)
                if inside and y + eh > cy + chh + 0.03:
                    issues.append(f"S{i} Text sprengt Karte '{nc}': "
                                  f"Karte endet {cy + chh:.2f}, Text bis {y + eh:.2f} "
                                  f"→ {txt[:45]!r}")
        # Karten gegen Karten
        for a in range(len(cards)):
            for b in range(a + 1, len(cards)):
                _, ax, ay, aw, ah = cards[a]
                _, bx, by, bw, bh = cards[b]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.03 and oy > 0.03:
                    contains = (ax <= bx and ay <= by and ax + aw >= bx + bw and ay + ah >= by + bh) or \
                               (bx <= ax and by <= ay and bx + bw >= ax + aw and by + bh >= ay + ah)
                    if not contains:
                        issues.append(f"S{i} Karten ueberlappen: {cards[a][0]} / {cards[b][0]} "
                                      f"({ox:.2f}×{oy:.2f})")
    for it in issues:
        print(it)
    print(f"--- {len(issues)} Befund(e)")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else
                  "SKP-Marktanalyse-Schulbegleitung.pptx"))
